"""
Shared python-docx styling helpers for the Team Cipher mentor documents.

Used by:
    generate_docx.py           -> Solution Design Note & Clarification Request
    generate_findings_docx.py  -> Dataset Analysis Findings

Kept in one place so both attachments are typographically identical and a styling
change never has to be made twice.
"""

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

# ---------------------------------------------------------------- palette
NAVY = RGBColor(0x1F, 0x33, 0x64)
ACCENT = RGBColor(0x00, 0x5A, 0xD2)
GREY = RGBColor(0x59, 0x59, 0x59)
CRIMSON = RGBColor(0xB0, 0x1C, 0x2E)

NAVY_HEX = "1F3364"
ACCENT_HEX = "005AD2"
LIGHT_HEX = "EEF2F9"
WARN_HEX = "FDF3F4"

BODY_PT = 10.5
TABLE_PT = 9.5


def new_document():
    """A document with the shared base styles and margins already applied."""
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(BODY_PT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    for sec in doc.sections:
        sec.top_margin = Inches(0.8)
        sec.bottom_margin = Inches(0.8)
        sec.left_margin = Inches(0.9)
        sec.right_margin = Inches(0.9)
    return doc


# ---------------------------------------------------------------- primitives
def shade(cell, hex_color):
    """Apply a solid background fill to a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), hex_color)
    tcPr.append(el)


def hrule(p):
    """Add a thin bottom rule to a paragraph."""
    pPr = p._p.get_or_add_pPr()
    bd = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:color"), "C6CEDA")
    bd.append(bottom)
    pPr.append(bd)


# ---------------------------------------------------------------- headings
def title(doc, text, size=19):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.color.rgb = NAVY
    p.paragraph_format.space_after = Pt(2)
    return p


def subtitle(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(11)
    r.font.color.rgb = ACCENT
    r.font.bold = True
    p.paragraph_format.space_after = Pt(10)
    hrule(p)
    return p


def h1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = NAVY
    hrule(p)
    return p


def h2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = ACCENT
    return p


# ---------------------------------------------------------------- text
def body(doc, text, italic=False, size=BODY_PT, space_after=6, bold=False, color=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.italic = italic
    r.font.bold = bold
    if color is not None:
        r.font.color.rgb = color
    p.paragraph_format.space_after = Pt(space_after)
    return p


def rich(doc, parts, space_after=6, size=BODY_PT):
    """
    A paragraph built from (text, bold) tuples, for inline emphasis.
    e.g. rich(doc, [("Result: ", True), ("5.27% precision", False)])
    """
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    for text, is_bold in parts:
        r = p.add_run(text)
        r.font.size = Pt(size)
        r.font.bold = is_bold
    return p


def bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Inches(0.28)
    if bold_prefix:
        r = p.add_run(bold_prefix)
        r.font.bold = True
        r.font.size = Pt(BODY_PT)
    r2 = p.add_run(text)
    r2.font.size = Pt(BODY_PT)
    return p


def numbered(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Inches(0.28)
    if bold_prefix:
        r = p.add_run(bold_prefix)
        r.font.bold = True
        r.font.size = Pt(BODY_PT)
    r2 = p.add_run(text)
    r2.font.size = Pt(BODY_PT)
    return p


def callout(doc, heading, text, tone="info"):
    """A single-cell shaded box for a finding that must not be skimmed past."""
    fill = WARN_HEX if tone == "warn" else LIGHT_HEX
    accent = CRIMSON if tone == "warn" else NAVY
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    cell = t.rows[0].cells[0]
    cell.text = ""
    shade(cell, fill)

    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(heading)
    r.font.bold = True
    r.font.size = Pt(10)
    r.font.color.rgb = accent

    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(2)
    r2 = p2.add_run(text)
    r2.font.size = Pt(10)
    body(doc, "", space_after=4)
    return t


# ---------------------------------------------------------------- tables
def table(doc, rows, headers, widths, bold_first_col=False):
    """A grid table with a navy header row."""
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER

    hdr = t.rows[0].cells
    for i, htext in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        r = p.add_run(htext)
        r.font.bold = True
        r.font.size = Pt(TABLE_PT)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade(hdr[i], NAVY_HEX)

    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(str(val))
            r.font.size = Pt(TABLE_PT)
            if bold_first_col and i == 0:
                r.font.bold = True

    for r_ in t.rows:
        for i, w in enumerate(widths):
            r_.cells[i].width = Inches(w)
    return t


def meta_table(doc, pairs):
    """The two-row label/value banner used at the top of both documents."""
    m = doc.add_table(rows=2, cols=len(pairs))
    m.style = "Table Grid"
    for i, (k, v) in enumerate(pairs):
        c1 = m.rows[0].cells[i]
        c1.text = ""
        r = c1.paragraphs[0].add_run(k)
        r.font.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade(c1, ACCENT_HEX)

        c2 = m.rows[1].cells[i]
        c2.text = ""
        r2 = c2.paragraphs[0].add_run(str(v))
        r2.font.size = Pt(9)
    return m


def footer(doc, text="Prepared by Team Cipher  |  Amity University  |  Novo Nordisk GBS Hackathon 2026"):
    p = doc.add_paragraph()
    hrule(p)
    p2 = doc.add_paragraph()
    r = p2.add_run(text)
    r.font.size = Pt(8.5)
    r.font.color.rgb = GREY
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p2
